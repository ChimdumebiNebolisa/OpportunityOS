"""Conservative local document extraction. Imported bytes are never executed."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

import filetype  # type: ignore[import-untyped]
from docx import Document
from pypdf import PdfReader

from opportunityos.config.settings import Settings
from opportunityos.infrastructure.storage import PrivateStorage, is_reparse_point


@dataclass(frozen=True)
class ExtractedDocument:
    private_path: Path
    content_hash: str
    mime_type: str
    text: str
    needs_vision: bool
    page_count: int | None


class DocumentExtractor:
    def __init__(self, settings: Settings, storage: PrivateStorage) -> None:
        self.settings = settings
        self.storage = storage

    def ingest(self, source_path: Path) -> ExtractedDocument:
        current = Path(source_path.absolute().anchor)
        for part in source_path.absolute().parts[1:]:
            current /= part
            if is_reparse_point(current):
                raise ValueError("Input cannot traverse a symlink or junction")
        source = source_path.resolve(strict=True)
        if not source.is_file():
            raise ValueError("Input must be a regular non-symlink file")
        size = source.stat().st_size
        if size <= 0 or size > self.settings.max_attachment_bytes:
            raise ValueError("Attachment is empty or exceeds the configured size limit")
        value = source.read_bytes()
        suffix = source.suffix.lower()
        mime = self._detect(value, suffix)
        if mime == "application/pdf":
            reader = PdfReader(io.BytesIO(value), strict=True)
            if len(reader.pages) > self.settings.max_pdf_pages:
                raise ValueError("PDF exceeds the configured page limit")
            text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
            needs_vision = not bool(text)
            page_count: int | None = len(reader.pages)
        elif mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            document = Document(io.BytesIO(value))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
            needs_vision = False
            page_count = None
        elif mime == "text/plain":
            text = value.decode("utf-8")
            needs_vision = False
            page_count = None
        elif mime.startswith("image/"):
            text = ""
            needs_vision = True
            page_count = None
        else:
            raise ValueError("Unsupported document type")
        private_path, content_hash = self.storage.store_bytes("attachments", value, suffix)
        return ExtractedDocument(private_path, content_hash, mime, text, needs_vision, page_count)

    @staticmethod
    def _detect(value: bytes, suffix: str) -> str:
        kind = filetype.guess(value)
        detected = kind.mime if kind else None
        if suffix == ".pdf":
            if not value.startswith(b"%PDF-") or detected not in {None, "application/pdf"}:
                raise ValueError("File extension and MIME signature do not match")
            return "application/pdf"
        if suffix == ".docx":
            try:
                with zipfile.ZipFile(io.BytesIO(value)) as archive:
                    entries = archive.infolist()
                    if len(entries) > 2_000 or sum(item.file_size for item in entries) > 50_000_000:
                        raise ValueError("DOCX archive exceeds the safe extraction limit")
                    names = {item.filename for item in entries}
                    valid = "[Content_Types].xml" in names and "word/document.xml" in names
                    if any(name.startswith("../") or name.startswith("/") for name in names):
                        raise ValueError("Unsafe archive path")
            except zipfile.BadZipFile as error:
                raise ValueError("File extension and MIME signature do not match") from error
            if not valid:
                raise ValueError("File extension and MIME signature do not match")
            return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if suffix in {".txt", ".md"}:
            if b"\x00" in value:
                raise ValueError("File extension and MIME signature do not match")
            try:
                value.decode("utf-8")
            except UnicodeDecodeError as error:
                raise ValueError("Text input must be UTF-8") from error
            return "text/plain"
        image_mimes = {"image/png", "image/jpeg", "image/webp"}
        if suffix in {".png", ".jpg", ".jpeg", ".webp"} and detected in image_mimes:
            return str(detected)
        raise ValueError("Unsupported or mismatched attachment type")
